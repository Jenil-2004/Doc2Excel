import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image, ImageDraw, ImageFont

def create_sample_assets():
    sample_dir = r"c:\CivilFileExcelOCR\samples"
    os.makedirs(sample_dir, exist_ok=True)

    # 1. Sample Excel Template A: Civil Construction Bill Template
    wb_civil = openpyxl.Workbook()
    ws_civil = wb_civil.active
    ws_civil.title = "Civil Measurements"

    # Title Banner (Row 1)
    ws_civil.merge_cells("A1:G1")
    title_cell = ws_civil["A1"]
    title_cell.value = "INFRASTRUCTURE PROJECT - CONTRACTOR BILLING SHEET"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_civil.row_dimensions[1].height = 30

    # Header Row (Row 2)
    headers_civil = ["Date", "Vendor", "Invoice Number", "Description", "Quantity", "Rate", "Amount"]
    ws_civil.append(headers_civil)
    ws_civil.row_dimensions[2].height = 24

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for col_num in range(1, len(headers_civil) + 1):
        cell = ws_civil.cell(row=2, column=col_num)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws_civil.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 18

    # Sample format row (Row 3)
    ws_civil.append(["", "", "", "", 0, 0.00, 0.00])
    ws_civil.cell(row=3, column=5).number_format = '#,##0'
    ws_civil.cell(row=3, column=6).number_format = '₹#,##0.00'
    ws_civil.cell(row=3, column=7).number_format = '₹#,##0.00'

    wb_civil.save(os.path.join(sample_dir, "Civil_Construction_Template.xlsx"))
    print("Saved Civil_Construction_Template.xlsx")

    # 2. Sample Excel Template B: General Expense & Restaurant/Retail Template
    wb_retail = openpyxl.Workbook()
    ws_retail = wb_retail.active
    ws_retail.title = "Expense Report"

    headers_retail = ["Bill Date", "Supplier", "Bill No", "Item Name", "Qty", "Unit Price", "Total Amount"]
    ws_retail.append(headers_retail)
    for col_num in range(1, len(headers_retail) + 1):
        cell = ws_retail.cell(row=1, column=col_num)
        cell.font = Font(name="Arial", size=10, bold=True, color="333333")
        cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_retail.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 16

    wb_retail.save(os.path.join(sample_dir, "General_Expense_Template.xlsx"))
    print("Saved General_Expense_Template.xlsx")

    # 3. Create Sample Bill Image: Construction Bill
    img_width, img_height = 800, 700
    img = Image.new('RGB', (img_width, img_height), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)

    # Simulated invoice text
    text_content = """APEX CIVIL CONSTRUCTIONS & BUILDERS LTD
GSTIN: 27AABCA1234F1Z8
Tax Invoice / Material Bill
------------------------------------------------------------
Invoice No: APEX-2026-904
Bill Date: 22/10/2026
Customer: Metro Infra Development Corp
Project Site: Sky Tower Phase 2

PARTICULARS / MATERIALS       QTY    UNIT   RATE     AMOUNT
------------------------------------------------------------
Ultratech Cement Grade 53     50     Bags   420.00   21000.00
River Sand Fine Quality       20     Tons   950.00   19000.00
Clay Burnt Red Bricks         5000   Nos    8.00     40000.00
TMT Steel Bars 12mm           15     Qntl   5800.00  87000.00
------------------------------------------------------------
Subtotal: Rs. 167000.00
Tax / GST (18%): Rs. 30060.00
Grand Total: Rs. 197060.00

Authorized Signatory
Apex Civil Constructions
"""
    y = 30
    for line in text_content.splitlines():
        draw.text((50, y), line, fill=(10, 10, 10))
        y += 24

    sample_img_path = os.path.join(sample_dir, "sample_construction_bill.png")
    img.save(sample_img_path)
    print(f"Saved {sample_img_path}")

if __name__ == "__main__":
    create_sample_assets()
